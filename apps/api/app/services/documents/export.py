"""合并稿导出：Markdown 与 docx。"""

import io
import re

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_LIST = re.compile(r"^\s*(?:[-*+])\s+(.*)$")
_MD_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_MD_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_MD_TABLE_DIVIDER = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def to_markdown(markdown: str) -> bytes:
    """直接以 UTF-8 输出 Markdown 文件内容。"""
    return markdown.encode("utf-8")


def _set_cjk_font(document) -> None:  # noqa: ANN001
    """让中文字体在 Word 里正确显示（默认西文字体不含中文）。"""
    try:
        from docx.oxml.ns import qn

        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except (KeyError, AttributeError):
        pass


def markdown_to_docx(markdown: str, title: str = "合并文稿") -> bytes:
    """把合并稿 Markdown 转成 docx（标题用 Word 内置样式，便于后续排版）。"""
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("服务器缺少 python-docx 依赖，无法导出 docx") from exc

    document = DocxDocument()
    _set_cjk_font(document)

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if _MD_TABLE_DIVIDER.match(line):
            continue

        heading = _MD_HEADING.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(heading.group(2).strip(), level=level)
            continue

        if _MD_TABLE_ROW.match(line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            document.add_paragraph(" | ".join(cells))
            continue

        quoted = _MD_QUOTE.match(line)
        if quoted:
            paragraph = document.add_paragraph(quoted.group(1).strip())
            paragraph.style = document.styles["Quote"]
            continue

        item = _MD_LIST.match(line)
        if item:
            document.add_paragraph(item.group(1).strip(), style="List Bullet")
            continue

        document.add_paragraph(line.strip())

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
