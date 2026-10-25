"""文档解析：docx / Markdown / 纯文本 → IR。

解析策略：
1. Markdown 的 `#` 标题最可靠，优先采用；
2. docx 先看段落样式（Heading N / 标题 N），再看大纲级别；
3. 都没有时，用中文编号习惯兜底（"一、"、"（二）"、"1.1" 等）。

标题自带的编号一律剥离，合并时统一重新编号。
"""

import io
import re

from app.services.documents.ir import Block, Node, ParsedDocument

SUPPORTED_SUFFIXES = {".docx", ".md", ".markdown", ".txt"}

_MAX_HEADING_LEN = 40


class ParseError(ValueError):
    """文件无法解析时抛出，附带可读原因。"""


# ---------------- 标题识别 ----------------

_HEADING_RULES: list[tuple[re.Pattern[str], int]] = [
    (re.compile(r"^第[一二三四五六七八九十百零]+[章篇]"), 1),
    (re.compile(r"^[一二三四五六七八九十]+[、.．]"), 1),
    (re.compile(r"^（[一二三四五六七八九十]+）"), 2),
    (re.compile(r"^\([一二三四五六七八九十]+\)"), 2),
    (re.compile(r"^\d+\.\d+\.\d+"), 3),
    (re.compile(r"^\d+\.\d+"), 2),
    (re.compile(r"^\d+[、.．]\s*"), 1),
]

_NUMBER_PREFIX = re.compile(
    r"^(第[一二三四五六七八九十百零]+[章篇]|[一二三四五六七八九十]+[、.．]|"
    r"（[一二三四五六七八九十]+）|\([一二三四五六七八九十]+\)|"
    r"\d+(?:\.\d+)*[、.．]?)\s*"
)


def strip_numbering(title: str) -> str:
    """去掉标题自带的编号，便于统一重新编号。"""
    stripped = _NUMBER_PREFIX.sub("", title).strip()
    return stripped or title.strip()


def detect_numbered_heading(text: str) -> int | None:
    """按中文编号习惯判断一行是否为标题，返回层级。"""
    stripped = text.strip()
    if not stripped or len(stripped) > _MAX_HEADING_LEN:
        return None
    if stripped.endswith(("。", "！", "？", "；", ";", "，", ",")):
        return None
    for pattern, level in _HEADING_RULES:
        if pattern.match(stripped):
            return level
    return None


# ---------------- Markdown / 纯文本 ----------------

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_LIST = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_MD_TABLE = re.compile(r"^\s*\|.*\|\s*$")
_MD_QUOTE = re.compile(r"^\s*>\s?")


def parse_text(text: str, filename: str, author: str = "") -> ParsedDocument:
    """解析 Markdown / 纯文本，按 # 的数量推断标题层级。

    代码块内的 # 不计入标题——这是纯文本解析最容易误判的地方。
    """
    nodes: list[Node] = []
    current = Node(level=0, title="", source=filename)
    nodes.append(current)
    in_code = False

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            current.blocks.append(Block("paragraph", line))
            continue
        if not line.strip():
            continue

        md = _MD_HEADING.match(line)
        if md:
            raw = md.group(2).strip()
            current = Node(
                level=len(md.group(1)),
                title=strip_numbering(raw),
                raw_title=raw,
                source=filename,
            )
            nodes.append(current)
            continue

        detected = detect_numbered_heading(line)
        if detected is not None:
            raw = line.strip()
            current = Node(
                level=detected, title=strip_numbering(raw), raw_title=raw, source=filename
            )
            nodes.append(current)
            continue

        if _MD_TABLE.match(line):
            current.blocks.append(Block("table", line.strip()))
        elif _MD_QUOTE.match(line):
            current.blocks.append(Block("quote", _MD_QUOTE.sub("", line).strip()))
        elif _MD_LIST.match(line):
            current.blocks.append(Block("list", _MD_LIST.sub("", line).strip()))
        else:
            current.blocks.append(Block("paragraph", line.strip()))

    nodes = [n for n in nodes if n.level >= 1 or n.blocks]
    return ParsedDocument(filename=filename, author=author, nodes=nodes)


# ---------------- docx ----------------

_DOCX_HEADING_STYLES = ("heading", "标题", "標題")


def _docx_heading_level(paragraph) -> int | None:  # noqa: ANN001 - python-docx 类型
    name = (getattr(paragraph.style, "name", "") or "").strip()
    lowered = name.lower()
    for prefix in _DOCX_HEADING_STYLES:
        if lowered.startswith(prefix):
            tail = lowered.replace(prefix, "").strip()
            if tail.isdigit():
                return max(1, min(6, int(tail)))
            return 1
    # 兜底：读取大纲级别
    try:
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.outlineLvl is not None:
            return max(1, min(6, int(ppr.outlineLvl.val) + 1))
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _table_to_lines(table) -> list[str]:  # noqa: ANN001
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append(cells)
    if not rows:
        return []
    width = max(len(r) for r in rows)
    out = ["| " + " | ".join(r + [""] * (width - len(r))) + " |" for r in rows]
    out.insert(1, "|" + "---|" * width)
    return out


def parse_docx(data: bytes, filename: str, author: str = "") -> ParsedDocument:
    """解析 docx。

    依赖缺失或文件损坏都抛 ParseError（而不是让原始异常冒出去），
    这样上层可以把多份文件的问题汇总成一条用户看得懂的提示。
    """
    try:
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - 依赖缺失时给出明确提示
        raise ParseError("服务器缺少 python-docx 依赖，无法解析 docx") from exc

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"无法读取 docx 文件（可能已损坏或不是真正的 docx）：{exc}") from exc

    nodes: list[Node] = []
    current = Node(level=0, title="", source=filename)
    nodes.append(current)

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _docx_heading_level(paragraph)
            if level is None:
                level = detect_numbered_heading(text)
            if level is not None:
                current = Node(
                    level=level, title=strip_numbering(text), raw_title=text, source=filename
                )
                nodes.append(current)
            else:
                current.blocks.append(Block("paragraph", text))
        elif tag == "tbl":
            table = Table(child, document)
            for line in _table_to_lines(table):
                current.blocks.append(Block("table", line))

    nodes = [n for n in nodes if n.level >= 1 or n.blocks]
    return ParsedDocument(filename=filename, author=author, nodes=nodes)


# ---------------- 统一入口 ----------------


def parse_document(filename: str, data: bytes, author: str = "") -> ParsedDocument:
    """按扩展名分发解析器。"""
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if suffix not in SUPPORTED_SUFFIXES:
        raise ParseError(
            f"不支持的文件类型 {suffix or '（无扩展名）'}，"
            f"目前支持：{', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    if not data:
        raise ParseError("文件内容为空")

    if suffix == ".docx":
        return parse_docx(data, filename, author)

    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = data.decode("gbk")
        except UnicodeDecodeError as exc:
            raise ParseError("文本编码无法识别（既不是 UTF-8 也不是 GBK）") from exc
    return parse_text(text, filename, author)
